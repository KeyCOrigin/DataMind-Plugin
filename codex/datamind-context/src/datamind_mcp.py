#!/usr/bin/env python3
"""MCP stdio server for the DataMind Codex plugin prototype."""

from __future__ import annotations

import datetime as dt
import hashlib
import json
import os
import pathlib
import re
import shutil
import subprocess
import sys
import traceback
from typing import Any, Callable

import memory_store


SERVER_NAME = "datamind-context"
SERVER_VERSION = "1.0.0"
DEFAULT_PROTOCOL_VERSION = "2025-06-18"
PLUGIN_ROOT = pathlib.Path(__file__).resolve().parents[1]


def bundled_repo_root() -> pathlib.Path | None:
    for candidate in (
        PLUGIN_ROOT / "vendor" / "datamind",
        PLUGIN_ROOT / "vendor" / "DataMind",
    ):
        if (candidate / "config.py").is_file() and (candidate / "core").is_dir():
            return candidate.resolve()
    return None


def resolve_repo_root() -> pathlib.Path:
    env_root = os.environ.get("DATAMIND_REPO_ROOT")
    marker_path = PLUGIN_ROOT / ".datamind-repo-root"
    if env_root:
        return pathlib.Path(env_root).expanduser().resolve()
    if marker_path.is_file():
        marker = marker_path.read_text(encoding="utf-8", errors="replace").splitlines()[0].strip()
        if marker:
            return pathlib.Path(marker).expanduser().resolve()
    bundled = bundled_repo_root()
    if bundled is not None:
        return bundled
    return PLUGIN_ROOT.parents[1].resolve()


REPO_ROOT = resolve_repo_root()
TASK_SCRIPT = PLUGIN_ROOT / "src" / "datamind_task.py"
CONFIG_DIR = pathlib.Path.home() / ".datamind-context"
PROFILES_PATH = CONFIG_DIR / "profiles.json"
MANAGED_PROFILES_DIR = CONFIG_DIR / "profiles"
INTERACTIONS_DIR = CONFIG_DIR / "interactions"
MEMORIES_DIR = CONFIG_DIR / "memories"
GENERATED_SKILLS_DIR = CONFIG_DIR / "skills"
GLOBAL_WIKI_DIR = CONFIG_DIR / "wiki"
QUERY_CONTEXT_DIR = CONFIG_DIR / "runtime" / "query-contexts"
MEMORY_DB_PATH = CONFIG_DIR / "memory.db"
DEFAULT_TIMEOUT_SECONDS = 900
GLOBAL_SESSION_ID = os.environ.get("DATAMIND_GLOBAL_SESSION_ID", "global")
TEXT_EXTENSIONS = {".txt", ".md", ".markdown"}
TABLE_EXTENSIONS = {".csv", ".tsv", ".xlsx", ".xls"}
DOC_EXTENSIONS = {".docx"}
PRESENTATION_EXTENSIONS = {".pptx"}
PDF_EXTENSIONS = {".pdf"}
MEDIA_EXTENSIONS = {".mp3", ".m4a", ".wav", ".aac", ".flac", ".mp4", ".mov", ".mkv", ".webm"}


class ToolError(Exception):
    """Expected tool-level error."""


def redact_secrets(value: str) -> str:
    return re.sub(r"sk-[A-Za-z0-9_-]+", "sk-<redacted>", value)


def eprint(message: str) -> None:
    print(f"[{SERVER_NAME}] {message}", file=sys.stderr, flush=True)


def expand_path(value: str | None) -> pathlib.Path | None:
    if not value:
        return None
    return pathlib.Path(os.path.expandvars(os.path.expanduser(value))).resolve()


def slugify(value: str) -> str:
    slug = re.sub(r"[^A-Za-z0-9_.-]+", "-", value).strip("-._").lower()
    return slug[:48] or "folder"


def derive_profile(folder_path: pathlib.Path) -> str:
    return slugify(folder_path.name)


def normalize_profile(value: str) -> str:
    profile = value.strip()
    if not re.fullmatch(r"[A-Za-z0-9_.-]{1,64}", profile):
        raise ToolError("profile must be 1-64 characters using only letters, numbers, underscore, dot, or dash.")
    if profile in {".", ".."}:
        raise ToolError("profile cannot be '.' or '..'.")
    return profile


def load_profiles() -> dict[str, Any]:
    if not PROFILES_PATH.exists():
        return {"profiles": {}}
    with PROFILES_PATH.open("r", encoding="utf-8") as handle:
        payload = json.load(handle)
    if not isinstance(payload, dict):
        return {"profiles": {}}
    profiles = payload.get("profiles")
    if not isinstance(profiles, dict):
        payload["profiles"] = {}
    return payload


def save_profiles(payload: dict[str, Any]) -> None:
    CONFIG_DIR.mkdir(parents=True, exist_ok=True)
    with PROFILES_PATH.open("w", encoding="utf-8") as handle:
        json.dump(payload, handle, ensure_ascii=False, indent=2)


def remember_profile(
    profile: str,
    folder_path: pathlib.Path,
    *,
    source_folder: pathlib.Path | None = None,
    managed: bool = False,
) -> dict[str, Any]:
    payload = load_profiles()
    existing = payload.setdefault("profiles", {}).get(profile, {})
    entry = {
        "profile": profile,
        "folder_path": str(folder_path),
        "data_dir": str(folder_path),
        "source_folder": str(source_folder or existing.get("source_folder") or folder_path),
        "managed": bool(managed or existing.get("managed", False)),
        "updated_at": dt.datetime.now(dt.timezone.utc).isoformat(),
    }
    payload["profiles"][profile] = entry
    payload["last_profile"] = profile
    save_profiles(payload)
    return entry


def resolve_profile_args(args: dict[str, Any], *, require_folder: bool = False) -> tuple[str, pathlib.Path | None]:
    raw_folder = args.get("folder_path") or args.get("data_dir")
    folder = expand_path(str(raw_folder)) if raw_folder else None
    if folder is not None and not folder.is_dir():
        raise ToolError(f"folder_path does not exist or is not a directory: {folder}")

    profile = str(args.get("profile") or "").strip()
    profiles_payload = load_profiles()
    profiles = profiles_payload.get("profiles", {})

    if not profile and folder is not None:
        profile = derive_profile(folder)
    elif profile and folder is None and profile in profiles:
        entry = profiles[profile]
        folder = expand_path(entry.get("data_dir") or entry.get("folder_path"))
    elif not profile:
        candidates = [
            profiles_payload.get("last_profile"),
            os.environ.get("DATA_PROFILE"),
            "default",
        ]
        for candidate in candidates:
            if candidate and candidate in profiles:
                profile = str(candidate)
                entry = profiles[profile]
                folder = expand_path(entry.get("data_dir") or entry.get("folder_path"))
                break
        if not profile:
            profile = os.environ.get("DATA_PROFILE", "default")

    profile = normalize_profile(profile)

    if require_folder and folder is None:
        raise ToolError("folder_path is required for this operation.")

    return profile, folder


def run_task(task: str, payload: dict[str, Any], timeout_seconds: int | None = None) -> dict[str, Any]:
    env = os.environ.copy()
    env.setdefault("PYTHONUNBUFFERED", "1")
    command = [sys.executable or "python3", str(TASK_SCRIPT), task]
    timeout = timeout_seconds or int(payload.get("timeout_seconds") or DEFAULT_TIMEOUT_SECONDS)

    completed = subprocess.run(
        command,
        input=json.dumps(payload, ensure_ascii=False),
        text=True,
        capture_output=True,
        cwd=str(REPO_ROOT),
        env=env,
        timeout=timeout,
    )
    stdout = completed.stdout.strip()
    stderr = completed.stderr.strip()
    try:
        result = json.loads(stdout.splitlines()[-1] if stdout else "{}")
    except json.JSONDecodeError as exc:
        raise ToolError(f"DataMind task returned invalid JSON: {exc}; stderr={stderr[-1200:]}")

    if completed.returncode != 0 or not result.get("ok"):
        message = result.get("error") or f"DataMind task failed with exit code {completed.returncode}"
        if stderr:
            message = f"{message}\n\nLogs:\n{stderr[-2000:]}"
        raise ToolError(redact_secrets(message))

    output = result.get("result") or {}
    if stderr:
        output["logs_tail"] = redact_secrets(stderr[-3000:])
    return output


def tool_result(payload: Any) -> dict[str, Any]:
    return {
        "content": [{"type": "text", "text": json.dumps(payload, ensure_ascii=False, indent=2)}],
        "structuredContent": payload,
    }


def error_result(message: str) -> dict[str, Any]:
    return {"content": [{"type": "text", "text": message}], "isError": True}


def utc_now() -> str:
    return dt.datetime.now(dt.timezone.utc).isoformat()


def session_slug(session_id: str) -> str:
    return slugify(session_id)[:64] or GLOBAL_SESSION_ID


def default_session_id(value: Any = None) -> str:
    return str(value or os.environ.get("DATAMIND_SESSION_ID") or GLOBAL_SESSION_ID)


def interaction_events_path(session_id: str) -> pathlib.Path:
    return INTERACTIONS_DIR / session_slug(session_id) / "events.jsonl"


def memory_file_path(session_id: str) -> pathlib.Path:
    return MEMORIES_DIR / session_slug(session_id) / "memory.md"


def skill_file_path(session_id: str) -> pathlib.Path:
    stamp = dt.datetime.now(dt.timezone.utc).strftime("%Y%m%dT%H%M%SZ")
    return GENERATED_SKILLS_DIR / session_slug(session_id) / f"{stamp}-skill.md"


def wiki_root(profile: str) -> pathlib.Path:
    return MANAGED_PROFILES_DIR / normalize_profile(profile) / "wiki"


def global_preferences_path() -> pathlib.Path:
    return GLOBAL_WIKI_DIR / "user" / "preferences.md"


def profile_preferences_path(profile: str) -> pathlib.Path:
    return wiki_root(profile) / "user" / "preferences.md"


def wiki_questions_dir(profile: str) -> pathlib.Path:
    return wiki_root(profile) / "questions"


def wiki_syntheses_dir(profile: str) -> pathlib.Path:
    return wiki_root(profile) / "syntheses"


def load_memory_context(session_id: str | None, max_chars: int = 12000) -> tuple[str, str | None]:
    if not session_id:
        return "", None
    path = memory_file_path(session_id)
    if not path.is_file():
        return "", str(path)
    return path.read_text(encoding="utf-8", errors="replace")[:max_chars], str(path)


def add_structured_memory_item(
    *,
    content: str,
    scope: str = "global",
    profile: str | None = None,
    session_id: str | None = None,
    kind: str | None = None,
    title: str | None = None,
    tags: Any = None,
    source: str = "",
    source_ref: str = "",
    metadata: dict[str, Any] | None = None,
    confidence: float = 0.7,
) -> dict[str, Any] | None:
    """Best-effort write into the structured memory store.

    The legacy Markdown/Wiki path must remain authoritative for compatibility,
    so structured-memory failures are logged and ignored.
    """
    text = redact_secrets(content.strip())
    if not text:
        return None
    metadata = metadata or {}
    resolved_kind = kind or memory_store.infer_kind(metadata)
    conn = None
    try:
        conn = memory_store.connect(MEMORY_DB_PATH)
        return memory_store.add_item(
            conn,
            content=text,
            scope=scope,
            profile=profile,
            session_id=session_id,
            kind=resolved_kind,
            title=title,
            tags=tags,
            source=source,
            source_ref=source_ref,
            metadata=metadata,
            confidence=confidence,
        )
    except Exception as exc:
        eprint(f"structured memory write skipped: {type(exc).__name__}: {exc}")
        return None
    finally:
        if conn is not None:
            conn.close()


def add_structured_note(
    content: str,
    *,
    profile: str | None = None,
    session_id: str | None = None,
    metadata: dict[str, Any] | None = None,
    source: str = "",
    source_ref: str = "",
    kind: str | None = None,
    title: str | None = None,
) -> dict[str, Any] | None:
    scope = "profile" if profile else "global"
    return add_structured_memory_item(
        content=content,
        scope=scope,
        profile=profile,
        session_id=session_id if scope == "session" else None,
        kind=kind,
        title=title,
        source=source,
        source_ref=source_ref,
        metadata=metadata,
    )


def structured_memory_context(
    *,
    query: str | None,
    profile: str,
    session_id: str | None,
    limit: int = 6,
) -> tuple[str, dict[str, Any]]:
    if not query or not query.strip():
        return "", {"used": False, "path": str(MEMORY_DB_PATH), "items": []}
    conn = None
    try:
        conn = memory_store.connect(MEMORY_DB_PATH)
        items = memory_store.search_items(
            conn,
            query=query,
            profile=profile,
            session_id=session_id,
            limit=limit,
        )
        text = memory_store.render_items_markdown(items)
        return text, {
            "used": bool(items),
            "path": str(MEMORY_DB_PATH),
            "item_count": len(items),
            "items": [
                {
                    "id": item.get("id"),
                    "kind": item.get("kind"),
                    "scope": item.get("scope"),
                    "profile": item.get("profile"),
                    "title": item.get("title"),
                }
                for item in items
            ],
        }
    except Exception as exc:
        eprint(f"structured memory search skipped: {type(exc).__name__}: {exc}")
        return "", {"used": False, "path": str(MEMORY_DB_PATH), "error": str(exc)}
    finally:
        if conn is not None:
            conn.close()


def structured_memory_stats() -> dict[str, Any]:
    conn = None
    try:
        conn = memory_store.connect(MEMORY_DB_PATH)
        return {"path": str(MEMORY_DB_PATH), **memory_store.stats(conn)}
    except Exception as exc:
        return {"path": str(MEMORY_DB_PATH), "error": str(exc)}
    finally:
        if conn is not None:
            conn.close()


def record_answer_trace(
    *,
    profile: str | None,
    session_id: str | None,
    mode: str,
    query: str,
    answer: str,
    sources: Any = None,
    query_context: Any = None,
    metadata: dict[str, Any] | None = None,
) -> dict[str, Any]:
    conn = None
    try:
        conn = memory_store.connect(MEMORY_DB_PATH)
        trace = memory_store.add_feedback_trace(
            conn,
            profile=profile,
            session_id=session_id,
            mode=mode,
            query=query,
            answer=answer,
            sources=sources,
            query_context=query_context,
            metadata=metadata,
        )
        return {
            "id": trace.get("id"),
            "path": str(MEMORY_DB_PATH),
            "profile": trace.get("profile"),
            "session_id": trace.get("session_id"),
            "mode": trace.get("mode"),
        }
    except Exception as exc:
        eprint(f"feedback trace write skipped: {type(exc).__name__}: {exc}")
        return {"path": str(MEMORY_DB_PATH), "written": False, "error": str(exc)}
    finally:
        if conn is not None:
            conn.close()


def parse_feedback_score(value: Any) -> int | None:
    if value is None:
        return None
    if isinstance(value, bool):
        return 1 if value else -1
    if isinstance(value, int) and value in (-1, 0, 1):
        return value
    raw = str(value).strip().lower()
    positive = {"+1", "1", "positive", "good", "success", "ok", "correct", "yes", "pass"}
    negative = {"-1", "negative", "bad", "fail", "wrong", "incorrect", "no", "correction"}
    neutral = {"0", "neutral", "unknown", "unclear"}
    if raw in positive:
        return 1
    if raw in negative:
        return -1
    if raw in neutral:
        return 0
    return None


def classify_feedback_signal(text: str, explicit_score: Any = None) -> tuple[int, str]:
    parsed = parse_feedback_score(explicit_score)
    if parsed is not None:
        return parsed, {1: "explicit_positive", 0: "explicit_neutral", -1: "explicit_negative"}[parsed]

    lower = text.lower()
    negative_patterns = [
        "不对", "不是", "错了", "错误", "重来", "重新", "修正", "纠正", "改成", "改一下",
        "漏了", "没提", "不要", "不需要", "误解", "再试", "幻觉", "编造",
        "wrong", "incorrect", "redo", "retry", "fix", "change", "instead", "missing",
        "not what", "not this", "hallucinat", "fabricat",
    ]
    positive_patterns = [
        "对", "可以", "很好", "不错", "准确", "继续", "收到", "谢谢", "没问题", "就是这样",
        "ok", "okay", "great", "thanks", "correct", "looks good", "works", "yes",
    ]
    if any(pattern in lower for pattern in negative_patterns):
        return -1, "correction_or_retry"
    if any(pattern in lower for pattern in positive_patterns):
        return 1, "positive_progress"
    return 0, "ambiguous_next_state"


def resolve_feedback_trace(args: dict[str, Any], conn) -> tuple[dict[str, Any] | None, str | None]:
    trace_id = str(args.get("trace_id") or "").strip()
    if trace_id:
        return memory_store.get_feedback_trace(conn, trace_id), trace_id

    profile_filter = None
    if args.get("profile") or args.get("folder_path"):
        profile_filter, _folder = resolve_profile_args(args)
    session_id = default_session_id(args.get("session_id")) if args.get("include_session", True) is not False else None
    trace = memory_store.latest_feedback_trace(conn, profile=profile_filter, session_id=session_id)
    if trace is None and profile_filter:
        trace = memory_store.latest_feedback_trace(conn, session_id=session_id)
    if trace is None:
        trace = memory_store.latest_feedback_trace(conn)
    return trace, trace.get("id") if trace else None


def build_feedback_hint(trace: dict[str, Any] | None, feedback_text: str, explicit_hint: str | None) -> str:
    if explicit_hint and explicit_hint.strip():
        return explicit_hint.strip()
    query = str((trace or {}).get("query") or "").strip()
    mode = str((trace or {}).get("mode") or "").strip()
    parts = ["DataMind next-state feedback hint."]
    if query:
        parts.append(f"Previous question: {query[:900]}")
    if mode:
        parts.append(f"Previous retrieval mode: {mode}.")
    parts.append(f"User correction or next-state feedback: {feedback_text.strip()[:1200]}")
    parts.append(
        "Future behavior: when a similar request appears, adjust retrieval mode, source selection, "
        "or answer framing according to this correction before responding."
    )
    return "\n".join(parts)


def record_feedback(args: dict[str, Any]) -> dict[str, Any]:
    feedback_text = str(args.get("feedback") or args.get("next_state") or args.get("content") or "").strip()
    if not feedback_text:
        raise ToolError("feedback, next_state, or content is required.")

    conn = None
    try:
        conn = memory_store.connect(MEMORY_DB_PATH)
        trace, trace_id = resolve_feedback_trace(args, conn)
        if args.get("trace_id") and trace is None:
            raise ToolError(f"Unknown feedback trace: {args.get('trace_id')}")

        score, label = classify_feedback_signal(feedback_text, args.get("score"))
        if args.get("label"):
            label = str(args.get("label"))
        profile = str(args.get("profile") or (trace or {}).get("profile") or "").strip() or None
        if profile:
            profile = normalize_profile(profile)
        session_id = str(args.get("session_id") or (trace or {}).get("session_id") or GLOBAL_SESSION_ID)
        hint = ""
        structured_item = None
        if score < 0 or str(args.get("hint") or "").strip():
            hint = build_feedback_hint(trace, feedback_text, str(args.get("hint") or "").strip())
        event = memory_store.add_feedback_event(
            conn,
            trace_id=trace_id,
            profile=profile,
            session_id=session_id,
            next_state=feedback_text,
            next_state_role=str(args.get("role") or args.get("next_state_role") or "user"),
            score=score,
            label=label,
            hint=hint,
            metadata={
                "source": args.get("source") or "datamind_record_feedback",
                "explicit_score": args.get("score"),
                "trace_found": trace is not None,
            },
        )
    except ToolError:
        raise
    except Exception as exc:
        raise ToolError(f"Feedback recording failed: {exc}") from exc
    finally:
        if conn is not None:
            conn.close()

    if hint:
        structured_item = add_structured_memory_item(
            content=hint,
            scope="profile" if profile else "global",
            profile=profile,
            session_id=None,
            kind="feedback_hint",
            title=f"Feedback hint: {str((trace or {}).get('query') or feedback_text)[:90]}",
            tags=["feedback", "next-state", str((trace or {}).get("mode") or "datamind")],
            source="datamind_record_feedback",
            source_ref=event.get("id") or trace_id or "",
            metadata={
                "trace_id": trace_id,
                "feedback_event_id": event.get("id"),
                "score": score,
                "label": label,
                "profile": profile,
            },
            confidence=0.85 if score < 0 else 0.75,
        )

    return tool_result(
        {
            "feedback_event": {
                "id": event.get("id"),
                "trace_id": trace_id,
                "score": score,
                "label": label,
                "profile": profile,
                "session_id": session_id,
            },
            "trace": {
                "id": trace_id,
                "query": (trace or {}).get("query"),
                "mode": (trace or {}).get("mode"),
                "found": trace is not None,
            },
            "structured_memory": (
                {
                    "id": structured_item.get("id"),
                    "kind": structured_item.get("kind"),
                    "scope": structured_item.get("scope"),
                    "path": str(MEMORY_DB_PATH),
                }
                if structured_item
                else {"path": str(MEMORY_DB_PATH), "written": False}
            ),
        }
    )


def feedback_status(args: dict[str, Any]) -> dict[str, Any]:
    profile = None
    if args.get("profile") or args.get("folder_path"):
        profile, _folder = resolve_profile_args(args)
    session_id = default_session_id(args.get("session_id")) if args.get("include_session", True) is not False else None
    limit = max(1, min(int(args.get("limit") or 10), 50))
    conn = None
    try:
        conn = memory_store.connect(MEMORY_DB_PATH)
        latest = memory_store.latest_feedback_trace(conn, profile=profile, session_id=session_id)
        if latest is None and profile:
            latest = memory_store.latest_feedback_trace(conn, session_id=session_id)
        return tool_result(
            {
                "memory_db_path": str(MEMORY_DB_PATH),
                "profile": profile,
                "session_id": session_id,
                "stats": memory_store.feedback_stats(conn, profile=profile, session_id=session_id),
                "latest_trace": (
                    {
                        "id": latest.get("id"),
                        "profile": latest.get("profile"),
                        "session_id": latest.get("session_id"),
                        "mode": latest.get("mode"),
                        "query": latest.get("query"),
                        "updated_at": latest.get("updated_at"),
                    }
                    if latest
                    else None
                ),
                "recent_feedback": memory_store.list_feedback_events(
                    conn,
                    profile=profile,
                    session_id=session_id,
                    limit=limit,
                    score=parse_feedback_score(args.get("score")),
                ),
            }
        )
    except Exception as exc:
        raise ToolError(f"Feedback status failed: {exc}") from exc
    finally:
        if conn is not None:
            conn.close()


def read_text_file(path: pathlib.Path, max_chars: int = 12000) -> str:
    if not path.is_file():
        return ""
    return path.read_text(encoding="utf-8", errors="replace")[:max_chars]


def append_markdown_log(path: pathlib.Path, line: str) -> None:
    path.parent.mkdir(parents=True, exist_ok=True)
    if not path.exists():
        title = path.stem.replace("-", " ").title()
        path.write_text(f"# {title}\n\n", encoding="utf-8")
    with path.open("a", encoding="utf-8") as handle:
        handle.write(line.rstrip() + "\n")


def append_wiki_index_item(index_path: pathlib.Path, title: str, description: str, item_path: pathlib.Path) -> None:
    index_path.parent.mkdir(parents=True, exist_ok=True)
    if not index_path.exists():
        index_path.write_text(f"# {title}\n\n{description}\n\n", encoding="utf-8")
    link = relative_link(index_path, item_path)
    with index_path.open("a", encoding="utf-8") as handle:
        handle.write(f"- {utc_now()}: [{item_path.stem}]({link})\n")


def wiki_page_name(rel_path: str, digest: str | None = None) -> str:
    base = slugify(pathlib.Path(rel_path).with_suffix("").as_posix().replace("/", "-"))[:80] or "source"
    suffix = f"-{digest[:8]}" if digest else ""
    return f"{base}{suffix}.md"


def wiki_note_name(text: str) -> str:
    stamp = dt.datetime.now(dt.timezone.utc).strftime("%Y%m%dT%H%M%SZ")
    digest = hashlib.sha1(text.encode("utf-8")).hexdigest()[:8]
    return f"{stamp}-{slugify(text)[:56]}-{digest}.md"


def relative_link(from_path: pathlib.Path, to_path: pathlib.Path) -> str:
    return os.path.relpath(to_path, start=from_path.parent).replace(os.sep, "/")


def source_preview(output_path: str | None, max_chars: int = 1800) -> str:
    if not output_path:
        return ""
    path = expand_path(output_path)
    if path is None or not path.is_file() or path.suffix.lower() not in {".txt", ".md", ".markdown", ".csv", ".tsv"}:
        return ""
    text = path.read_text(encoding="utf-8", errors="replace").strip()
    return text[:max_chars]


def ensure_preferences_doc(path: pathlib.Path, title: str) -> None:
    if not path.exists():
        path.parent.mkdir(parents=True, exist_ok=True)
        path.write_text(f"# {title}\n\n", encoding="utf-8")


def append_preference_to_wiki(
    content: str,
    *,
    profile: str | None = None,
    metadata: dict[str, Any] | None = None,
) -> dict[str, Any]:
    text = redact_secrets(content.strip())
    if not text:
        return {"written": []}
    metadata = metadata or {}
    stamp = utc_now()
    bullet = f"- {stamp}: {text}"
    if metadata:
        compact_meta = json.dumps(metadata, ensure_ascii=False, sort_keys=True)
        bullet += f"\n  - metadata: `{compact_meta}`"
    written = []

    global_path = global_preferences_path()
    ensure_preferences_doc(global_path, "Global User Preferences")
    append_markdown_log(global_path, bullet)
    written.append(str(global_path))

    target_profile = profile
    if not target_profile:
        last_profile = load_profiles().get("last_profile")
        target_profile = str(last_profile) if last_profile else None
    if target_profile:
        pref_path = profile_preferences_path(target_profile)
        ensure_preferences_doc(pref_path, f"User Preferences: {target_profile}")
        append_markdown_log(pref_path, bullet)
        written.append(str(pref_path))
    return {"written": written}


def normalize_messages(args: dict[str, Any]) -> list[dict[str, Any]]:
    raw_messages = args.get("messages")
    if raw_messages is None:
        role = str(args.get("role") or "user")
        content = str(args.get("content") or "").strip()
        if not content:
            raise ToolError("content or messages is required.")
        raw_messages = [{"role": role, "content": content}]
    if not isinstance(raw_messages, list):
        raise ToolError("messages must be an array.")

    messages = []
    for item in raw_messages:
        if not isinstance(item, dict):
            raise ToolError("Each message must be an object.")
        role = str(item.get("role") or "user")
        content = str(item.get("content") or "").strip()
        if not content:
            continue
        messages.append(
            {
                "role": role,
                "content": content,
                "timestamp": item.get("timestamp") or utc_now(),
                "metadata": item.get("metadata") or {},
            }
        )
    if not messages:
        raise ToolError("No non-empty messages supplied.")
    return messages


def append_interaction_event(args: dict[str, Any]) -> dict[str, Any]:
    session_id = default_session_id(args.get("session_id"))
    path = interaction_events_path(session_id)
    path.parent.mkdir(parents=True, exist_ok=True)
    event = {
        "session_id": session_id,
        "recorded_at": utc_now(),
        "source": args.get("source") or "codex",
        "profile": args.get("profile"),
        "messages": normalize_messages(args),
        "metadata": args.get("metadata") or {},
    }
    with path.open("a", encoding="utf-8") as handle:
        handle.write(json.dumps(event, ensure_ascii=False) + "\n")
    return {
        "session_id": session_id,
        "events_path": str(path),
        "messages_recorded": len(event["messages"]),
        "recorded_at": event["recorded_at"],
    }


def auto_record_global(content: str, *, profile: str | None = None, metadata: dict[str, Any] | None = None) -> None:
    try:
        output = append_interaction_event(
            {
                "session_id": GLOBAL_SESSION_ID,
                "profile": profile,
                "source": "codex-auto",
                "messages": [{"role": "system", "content": content}],
                "metadata": {"auto_recorded": True, **(metadata or {})},
            }
        )
        add_structured_memory_item(
            content=content,
            scope="global",
            profile=profile,
            session_id=GLOBAL_SESSION_ID,
            kind=memory_store.infer_kind(metadata, default="memory"),
            source="codex-auto",
            source_ref=output.get("events_path", ""),
            metadata={"auto_recorded": True, **(metadata or {})},
        )
    except Exception as exc:
        eprint(f"auto memory record skipped: {type(exc).__name__}: {exc}")


def record_interaction(args: dict[str, Any]) -> dict[str, Any]:
    output = append_interaction_event(args)
    return tool_result(
        {
            **output,
            "default_session": GLOBAL_SESSION_ID,
        }
    )


def list_memories(_: dict[str, Any]) -> dict[str, Any]:
    sessions = []
    session_ids = set()
    for base in (INTERACTIONS_DIR, MEMORIES_DIR, GENERATED_SKILLS_DIR):
        if base.exists():
            session_ids.update(path.name for path in base.iterdir() if path.is_dir())
    for sid in sorted(session_ids):
        events_path = INTERACTIONS_DIR / sid / "events.jsonl"
        memory_path = MEMORIES_DIR / sid / "memory.md"
        skills_dir = GENERATED_SKILLS_DIR / sid
        event_count = 0
        if events_path.is_file():
            with events_path.open("r", encoding="utf-8") as handle:
                event_count = sum(1 for line in handle if line.strip())
        skills = sorted(str(path) for path in skills_dir.glob("*.md")) if skills_dir.is_dir() else []
        sessions.append(
            {
                "session_id": sid,
                "events_path": str(events_path),
                "event_count": event_count,
                "memory_path": str(memory_path) if memory_path.is_file() else None,
                "skills": skills,
            }
        )
    return tool_result(
        {
            "sessions": sessions,
            "config_dir": str(CONFIG_DIR),
            "structured_memory": structured_memory_stats(),
        }
    )


def search_structured_memory(args: dict[str, Any]) -> dict[str, Any]:
    query = str(args.get("query") or "").strip()
    if not query:
        raise ToolError("query is required.")
    profile, _folder = resolve_profile_args(args) if args.get("profile") or args.get("folder_path") else (load_profiles().get("last_profile") or os.environ.get("DATA_PROFILE", "default"), None)
    profile = normalize_profile(str(profile))
    session_id = default_session_id(args.get("session_id")) if args.get("include_session", True) is not False else None
    limit = int(args.get("limit") or 10)
    conn = None
    try:
        conn = memory_store.connect(MEMORY_DB_PATH)
        items = memory_store.search_items(
            conn,
            query=query,
            profile=profile,
            session_id=session_id,
            limit=max(1, min(limit, 50)),
        )
        return tool_result(
            {
                "query": query,
                "profile": profile,
                "session_id": session_id,
                "memory_db_path": str(MEMORY_DB_PATH),
                "count": len(items),
                "items": items,
            }
        )
    except Exception as exc:
        raise ToolError(f"Structured memory search failed: {exc}") from exc
    finally:
        if conn is not None:
            conn.close()


def file_digest(path: pathlib.Path) -> str:
    h = hashlib.sha1()
    with path.open("rb") as handle:
        for chunk in iter(lambda: handle.read(1024 * 1024), b""):
            h.update(chunk)
    return h.hexdigest()


def safe_materialized_name(rel_path: pathlib.Path, digest: str, suffix: str) -> str:
    stem = slugify("-".join(rel_path.with_suffix("").parts))[:80] or "document"
    return f"{stem}-{digest[:10]}{suffix}"


def ignored_source_reason(path: pathlib.Path) -> str | None:
    if path.name.startswith("~$"):
        return "office temporary lock file"
    return None


def remove_managed_output(path_value: Any, target: pathlib.Path) -> bool:
    if not path_value:
        return False
    try:
        output_path = expand_path(str(path_value))
        if output_path is None or not output_path.is_file():
            return False
        output_path.relative_to(target)
    except Exception:
        return False
    output_path.unlink()
    return True


def markdown_header(source_path: pathlib.Path, rel_path: pathlib.Path) -> str:
    return f"<!-- source: {source_path} -->\n<!-- relative_path: {rel_path.as_posix()} -->\n\n"


def write_markdown(path: pathlib.Path, title: str, body: str, source_path: pathlib.Path, rel_path: pathlib.Path) -> None:
    path.parent.mkdir(parents=True, exist_ok=True)
    path.write_text(
        markdown_header(source_path, rel_path) + f"# {title}\n\n{body.strip()}\n",
        encoding="utf-8",
    )


def extract_table(path: pathlib.Path, rel_path: pathlib.Path, target: pathlib.Path, max_rows: int) -> dict[str, Any]:
    import pandas as pd

    digest = file_digest(path)
    out_path = target / safe_materialized_name(rel_path, digest, ".md")
    ext = path.suffix.lower()
    parts = []
    if ext in {".csv", ".tsv"}:
        sep = "\t" if ext == ".tsv" else ","
        df = pd.read_csv(path, sep=sep)
        parts.append(f"## {path.name}\n\nRows: {len(df)}, Columns: {len(df.columns)}\n\n")
        parts.append("```csv\n" + df.head(max_rows).to_csv(index=False) + "```")
    else:
        sheets = pd.read_excel(path, sheet_name=None)
        for sheet_name, df in sheets.items():
            parts.append(f"## Sheet: {sheet_name}\n\nRows: {len(df)}, Columns: {len(df.columns)}\n\n")
            parts.append("```csv\n" + df.head(max_rows).to_csv(index=False) + "```")
            parts.append("\n")
    write_markdown(out_path, rel_path.as_posix(), "\n\n".join(parts), path, rel_path)
    return {"status": "extracted", "output": str(out_path), "kind": "table", "sha1": digest}


def extract_docx(path: pathlib.Path, rel_path: pathlib.Path, target: pathlib.Path) -> dict[str, Any]:
    from docx import Document

    digest = file_digest(path)
    out_path = target / safe_materialized_name(rel_path, digest, ".md")
    doc = Document(str(path))
    parts = [p.text.strip() for p in doc.paragraphs if p.text.strip()]
    for table in doc.tables:
        rows = []
        for row in table.rows:
            rows.append(" | ".join(cell.text.strip().replace("\n", " ") for cell in row.cells))
        if rows:
            parts.append("\n".join(rows))
    write_markdown(out_path, rel_path.as_posix(), "\n\n".join(parts), path, rel_path)
    return {"status": "extracted", "output": str(out_path), "kind": "docx", "sha1": digest}


def extract_pptx(path: pathlib.Path, rel_path: pathlib.Path, target: pathlib.Path) -> dict[str, Any]:
    from pptx import Presentation

    digest = file_digest(path)
    out_path = target / safe_materialized_name(rel_path, digest, ".md")
    prs = Presentation(str(path))
    parts = []
    for idx, slide in enumerate(prs.slides, 1):
        texts = []
        for shape in slide.shapes:
            if hasattr(shape, "text"):
                text = str(shape.text).strip()
                if text:
                    texts.append(text)
        if texts:
            parts.append(f"## Slide {idx}\n\n" + "\n\n".join(texts))
    write_markdown(out_path, rel_path.as_posix(), "\n\n".join(parts), path, rel_path)
    return {"status": "extracted", "output": str(out_path), "kind": "pptx", "sha1": digest}


def parse_env_file() -> dict[str, str]:
    env_path = REPO_ROOT / ".env"
    values: dict[str, str] = {}
    if not env_path.exists():
        return values
    with env_path.open("r", encoding="utf-8") as handle:
        for line in handle:
            line = line.strip()
            if not line or line.startswith("#") or "=" not in line:
                continue
            key, value = line.split("=", 1)
            values[key.strip()] = value.strip().strip('"').strip("'")
    return values


def transcribe_media(path: pathlib.Path) -> str:
    from openai import OpenAI

    env = parse_env_file()
    api_key = env.get("LLM_API_KEY") or env.get("OPENAI_API_KEY")
    api_base = env.get("LLM_API_BASE") or "https://api.openai.com/v1"
    model = env.get("AUDIO_TRANSCRIPTION_MODEL") or "whisper-1"
    if not api_key:
        raise ToolError("LLM_API_KEY is required for media transcription.")
    client = OpenAI(api_key=api_key, base_url=api_base)
    with path.open("rb") as handle:
        result = client.audio.transcriptions.create(model=model, file=handle)
    return getattr(result, "text", str(result))


def materialize_folder(args: dict[str, Any]) -> tuple[str, pathlib.Path, pathlib.Path, dict[str, Any]]:
    source = expand_path(str(args.get("folder_path") or ""))
    if source is None or not source.is_dir():
        raise ToolError(f"folder_path does not exist or is not a directory: {source}")
    profile = normalize_profile(str(args.get("profile") or derive_profile(source)))
    target = MANAGED_PROFILES_DIR / profile / "data"
    manifest_path = MANAGED_PROFILES_DIR / profile / "manifest.json"
    max_rows = int(args.get("max_table_rows") or 500)
    clean = args.get("clean", True) is not False
    transcribe = args.get("transcribe_media", False) is True
    old_manifest: dict[str, Any] = {}
    if manifest_path.is_file() and not clean:
        try:
            old_manifest = json.loads(manifest_path.read_text(encoding="utf-8"))
        except Exception:
            old_manifest = {}
    old_by_rel = {
        str(item.get("relative_path")): item
        for item in old_manifest.get("files", [])
        if isinstance(item, dict) and item.get("relative_path")
    }
    ready_statuses = {"copied", "extracted", "transcribed"}
    changes: dict[str, list[Any]] = {
        "added": [],
        "changed": [],
        "removed": [],
        "unchanged": [],
        "ignored": [],
        "indexed_added": [],
        "indexed_changed": [],
        "indexed_removed": [],
    }

    if clean and target.exists():
        shutil.rmtree(target)
    target.mkdir(parents=True, exist_ok=True)

    files = []
    seen_rel_paths: set[str] = set()
    for path in sorted(source.rglob("*")):
        if not path.is_file() or any(part.startswith(".") for part in path.relative_to(source).parts):
            continue
        rel = path.relative_to(source)
        rel_key = rel.as_posix()
        ignore_reason = ignored_source_reason(path)
        if ignore_reason:
            changes["ignored"].append({"relative_path": rel_key, "reason": ignore_reason})
            continue
        seen_rel_paths.add(rel_key)
        ext = path.suffix.lower()
        digest = file_digest(path)
        old_item = old_by_rel.get(rel_key)
        item = {
            "source": str(path),
            "relative_path": rel_key,
            "size": path.stat().st_size,
            "mtime": path.stat().st_mtime,
            "extension": ext,
            "sha1": digest,
        }

        old_output = old_item.get("output") if isinstance(old_item, dict) else None
        old_status = old_item.get("status") if isinstance(old_item, dict) else None
        output_reusable = not old_output or pathlib.Path(str(old_output)).is_file()
        if old_item and old_item.get("sha1") == digest and output_reusable and not clean:
            reused = {**old_item, **item, "change_status": "unchanged", "reused": True}
            files.append(reused)
            changes["unchanged"].append(rel_key)
            continue

        change_kind = "changed" if old_item else "added"
        try:
            if ext in TEXT_EXTENSIONS:
                out_path = target / safe_materialized_name(rel, digest, ext)
                out_path.parent.mkdir(parents=True, exist_ok=True)
                shutil.copy2(path, out_path)
                item.update({"status": "copied", "output": str(out_path), "kind": "text", "sha1": digest})
            elif ext in PDF_EXTENSIONS:
                out_path = target / safe_materialized_name(rel, digest, ext)
                shutil.copy2(path, out_path)
                item.update({"status": "copied", "output": str(out_path), "kind": "pdf", "sha1": digest})
            elif ext in TABLE_EXTENSIONS:
                item.update(extract_table(path, rel, target, max_rows))
            elif ext in DOC_EXTENSIONS:
                item.update(extract_docx(path, rel, target))
            elif ext in PRESENTATION_EXTENSIONS:
                item.update(extract_pptx(path, rel, target))
            elif ext in MEDIA_EXTENSIONS:
                if transcribe:
                    digest = file_digest(path)
                    out_path = target / safe_materialized_name(rel, digest, ".md")
                    transcript = transcribe_media(path)
                    write_markdown(out_path, rel.as_posix(), f"## Transcript\n\n{transcript}", path, rel)
                    item.update({"status": "transcribed", "output": str(out_path), "kind": "media", "sha1": digest})
                else:
                    item.update({"status": "skipped", "kind": "media", "reason": "transcribe_media=false", "sha1": digest})
            else:
                item.update({"status": "skipped", "kind": "unsupported", "reason": "unsupported extension", "sha1": digest})
        except Exception as exc:
            item.update({"status": "error", "sha1": digest, "error": redact_secrets(f"{type(exc).__name__}: {exc}")})
        item["change_status"] = change_kind
        changes[change_kind].append(rel_key)
        if item.get("status") in ready_statuses:
            changes[f"indexed_{change_kind}"].append(rel_key)
        if old_output and old_output != item.get("output") and remove_managed_output(old_output, target):
            if old_status in ready_statuses:
                changes["indexed_removed"].append(rel_key)
        files.append(item)

    for rel_key, old_item in old_by_rel.items():
        if rel_key in seen_rel_paths:
            continue
        changes["removed"].append(rel_key)
        if remove_managed_output(old_item.get("output"), target) and old_item.get("status") in ready_statuses:
            changes["indexed_removed"].append(rel_key)

    manifest = {
        "profile": profile,
        "source_folder": str(source),
        "data_dir": str(target),
        "generated_at": utc_now(),
        "files": files,
        "changes": changes,
        "counts": {
            "total": len(files),
            "ready": sum(1 for item in files if item.get("status") in {"copied", "extracted", "transcribed"}),
            "skipped": sum(1 for item in files if item.get("status") == "skipped"),
            "errors": sum(1 for item in files if item.get("status") == "error"),
            "ignored": len(changes["ignored"]),
        },
    }
    manifest_path.parent.mkdir(parents=True, exist_ok=True)
    manifest_path.write_text(json.dumps(manifest, ensure_ascii=False, indent=2), encoding="utf-8")
    return profile, source, target, manifest


def write_profile_wiki(profile: str, manifest: dict[str, Any], profile_entry: dict[str, Any]) -> dict[str, Any]:
    """Create or refresh the lightweight Markdown Wiki for a profile."""
    root = wiki_root(profile)
    sources_dir = root / "sources"
    questions_dir = wiki_questions_dir(profile)
    syntheses_dir = wiki_syntheses_dir(profile)
    user_dir = root / "user"
    root.mkdir(parents=True, exist_ok=True)
    sources_dir.mkdir(parents=True, exist_ok=True)
    questions_dir.mkdir(parents=True, exist_ok=True)
    syntheses_dir.mkdir(parents=True, exist_ok=True)
    user_dir.mkdir(parents=True, exist_ok=True)

    counts = manifest.get("counts", {}) if isinstance(manifest, dict) else {}
    files = manifest.get("files", []) if isinstance(manifest, dict) else []
    source_pages = []

    for item in files:
        if not isinstance(item, dict):
            continue
        rel_path = str(item.get("relative_path") or pathlib.Path(str(item.get("source") or "source")).name)
        digest = str(item.get("sha1") or "") or None
        page_path = sources_dir / wiki_page_name(rel_path, digest)
        preview = source_preview(item.get("output"))
        page = [
            f"# Source: {rel_path}",
            "",
            "## Provenance",
            "",
            f"- Source file: `{item.get('source')}`",
            f"- Managed output: `{item.get('output')}`",
            f"- Status: `{item.get('status')}`",
            f"- Kind: `{item.get('kind')}`",
            f"- Extension: `{item.get('extension')}`",
            f"- SHA1: `{item.get('sha1')}`",
            "",
            "## Notes",
            "",
            "- This page is generated by DataMind's lightweight LLM Wiki layer.",
            "- Add durable human or LLM synthesis below this section if needed.",
        ]
        if item.get("reason"):
            page.extend(["", "## Skip Reason", "", str(item.get("reason"))])
        if item.get("error"):
            page.extend(["", "## Error", "", str(item.get("error"))])
        if preview:
            page.extend(["", "## Extract Preview", "", "```text", preview, "```"])
        page_path.write_text("\n".join(page).rstrip() + "\n", encoding="utf-8")
        source_pages.append({"relative_path": rel_path, "page": page_path, "status": item.get("status"), "kind": item.get("kind")})

    schema_path = root / "schema.md"
    schema_path.write_text(
        "\n".join(
            [
                "# Wiki Schema",
                "",
                "- `index.md`: profile overview and source map.",
                "- `sources/`: one provenance page per source file or extracted file.",
                "- `questions/`: high-value questions captured from Codex conversations.",
                "- `syntheses/`: reusable answer-derived synthesis notes.",
                "- `user/preferences.md`: durable user preferences relevant to this profile.",
                "- `log.md`: chronological changes to this Wiki layer.",
                "",
                "The Wiki is a cumulative Markdown layer over raw DataMind extracts. RAG and GraphRAG remain the fact retrieval layer.",
            ]
        )
        + "\n",
        encoding="utf-8",
    )

    preferences_path = user_dir / "preferences.md"
    ensure_preferences_doc(preferences_path, f"User Preferences: {profile}")
    if not (questions_dir / "index.md").exists():
        (questions_dir / "index.md").write_text("# Saved Questions\n\nHigh-value DataMind questions captured from Codex conversations.\n", encoding="utf-8")
    if not (syntheses_dir / "index.md").exists():
        (syntheses_dir / "index.md").write_text("# Syntheses\n\nReusable answer-derived synthesis notes.\n", encoding="utf-8")

    index_path = root / "index.md"
    lines = [
        f"# DataMind Wiki: {profile}",
        "",
        "## Profile",
        "",
        f"- Profile: `{profile}`",
        f"- Source folder: `{profile_entry.get('source_folder') or manifest.get('source_folder')}`",
        f"- Managed data dir: `{profile_entry.get('data_dir') or manifest.get('data_dir')}`",
        f"- Updated at: `{utc_now()}`",
        "",
        "## File Counts",
        "",
        f"- Total: `{counts.get('total', 0)}`",
        f"- Ready: `{counts.get('ready', 0)}`",
        f"- Skipped: `{counts.get('skipped', 0)}`",
        f"- Errors: `{counts.get('errors', 0)}`",
        "",
        "## Source Pages",
        "",
    ]
    if source_pages:
        for page in source_pages:
            link = relative_link(index_path, page["page"])
            lines.append(f"- [{page['relative_path']}]({link}) - `{page.get('status')}` / `{page.get('kind')}`")
    else:
        lines.append("- No source pages generated yet.")
    lines.extend(
        [
            "",
            "## User Context",
            "",
            f"- [Profile preferences]({relative_link(index_path, preferences_path)})",
            f"- [Saved questions]({relative_link(index_path, questions_dir / 'index.md')})",
            f"- [Syntheses]({relative_link(index_path, syntheses_dir / 'index.md')})",
            f"- [Wiki schema]({relative_link(index_path, schema_path)})",
            "- Global preferences are stored separately under `~/.datamind-context/wiki/user/preferences.md`.",
            "",
            "## Query Behavior",
            "",
            "DataMind queries use this Wiki as a high-level accumulated context, then use RAG or GraphRAG for source-grounded details.",
        ]
    )
    index_path.write_text("\n".join(lines).rstrip() + "\n", encoding="utf-8")

    log_path = root / "log.md"
    append_markdown_log(
        log_path,
        f"- {utc_now()}: Refreshed lightweight Wiki for profile `{profile}` with {len(source_pages)} source pages.",
    )

    return {
        "wiki_dir": str(root),
        "index_path": str(index_path),
        "log_path": str(log_path),
        "schema_path": str(schema_path),
        "preferences_path": str(preferences_path),
        "source_pages": len(source_pages),
    }


def load_wiki_context(profile: str, max_chars: int = 12000) -> tuple[str, dict[str, Any]]:
    root = wiki_root(profile)
    paths = [
        root / "index.md",
        root / "syntheses" / "index.md",
        root / "questions" / "index.md",
        root / "user" / "preferences.md",
        GLOBAL_WIKI_DIR / "user" / "preferences.md",
        root / "log.md",
    ]
    syntheses_dir = root / "syntheses"
    if syntheses_dir.is_dir():
        recent_syntheses = sorted(
            [path for path in syntheses_dir.glob("*.md") if path.name != "index.md"],
            key=lambda path: path.stat().st_mtime,
            reverse=True,
        )[:3]
        paths.extend(recent_syntheses)
    sections = []
    used_paths = []
    per_file = max(1200, max_chars // max(len(paths), 1))
    for path in paths:
        text = read_text_file(path, per_file).strip()
        if text:
            sections.append(f"<!-- DataMind Wiki: {path} -->\n{text}")
            used_paths.append(str(path))
    joined = "\n\n".join(sections)
    return joined[:max_chars], {"paths": used_paths, "wiki_dir": str(root), "used": bool(sections)}


def write_query_context(
    profile: str,
    session_id: str | None,
    include_memory: bool = True,
    *,
    query: str | None = None,
) -> tuple[str, str | None, dict[str, Any]]:
    wiki_text, wiki_meta = load_wiki_context(profile)
    structured_text = ""
    structured_meta: dict[str, Any] = {"used": False, "path": str(MEMORY_DB_PATH), "items": []}
    memory_text = ""
    memory_path = None
    if include_memory and session_id:
        structured_text, structured_meta = structured_memory_context(
            query=query,
            profile=profile,
            session_id=session_id,
        )
        memory_text, memory_path = load_memory_context(session_id)

    sections = []
    if wiki_text:
        sections.append("# DataMind LLM Wiki Context\n\n" + wiki_text)
    if structured_text:
        sections.append(structured_text)
    if memory_text:
        sections.append(f"# DataMind Session Memory: {session_id}\n\n" + memory_text)
    if not sections:
        return "", None, {"wiki": wiki_meta, "structured_memory": structured_meta, "memory_path": memory_path, "used": False}

    text = "\n\n---\n\n".join(sections)
    digest = hashlib.sha1(f"{profile}:{session_id}:{text}".encode("utf-8")).hexdigest()[:12]
    context_path = QUERY_CONTEXT_DIR / f"{profile}-{session_slug(session_id or GLOBAL_SESSION_ID)}-{digest}.md"
    context_path.parent.mkdir(parents=True, exist_ok=True)
    context_path.write_text(text, encoding="utf-8")
    return text, str(context_path), {
        "used": True,
        "path": str(context_path),
        "wiki": wiki_meta,
        "structured_memory": structured_meta,
        "memory_path": memory_path if memory_text else None,
        "memory_used": bool(memory_text),
    }


def should_capture_qa(question: str, answer: str, mode: str, capture_mode: str) -> tuple[bool, str]:
    capture_mode = capture_mode.lower().strip()
    if capture_mode == "never":
        return False, "disabled"
    if capture_mode == "always":
        return True, "forced"
    if capture_mode != "auto":
        raise ToolError("wiki_capture must be one of: auto, always, never.")

    high_value_keywords = [
        "总结", "结论", "建议", "方案", "计划", "风险", "决策", "比较", "分析",
        "关键", "下一步", "复盘", "设计", "策略", "路线", "框架", "原则",
        "summary", "conclusion", "recommendation", "plan", "risk", "decision",
        "compare", "analysis", "strategy", "design", "framework",
    ]
    combined = f"{question}\n{answer}".lower()
    if mode == "both":
        return True, "multi_mode_answer"
    if any(keyword.lower() in combined for keyword in high_value_keywords) and len(answer.strip()) >= 160:
        return True, "high_value_keyword"
    if len(answer.strip()) >= 900 and len(question.strip()) >= 12:
        return True, "substantial_answer"
    return False, "not_high_value"


def render_sources_for_markdown(sources: list[dict[str, Any]], limit: int = 8) -> str:
    if not sources:
        return "- No sources returned."
    lines = []
    for idx, source in enumerate(sources[:limit], 1):
        metadata = source.get("metadata") if isinstance(source, dict) else {}
        if not isinstance(metadata, dict):
            metadata = {}
        label = metadata.get("file_path") or metadata.get("filename") or metadata.get("source") or metadata.get("relative_path") or f"source-{idx}"
        score = source.get("score") if isinstance(source, dict) else None
        preview = str(source.get("text_preview") or "").strip() if isinstance(source, dict) else ""
        lines.append(f"- {idx}. `{label}`" + (f" score=`{score}`" if score is not None else ""))
        if preview:
            lines.append(f"  - preview: {preview[:350].replace(chr(10), ' ')}")
    return "\n".join(lines)


def save_qa_to_wiki(
    *,
    profile: str,
    question: str,
    answer: str,
    mode: str,
    sources: list[dict[str, Any]],
    query_context: dict[str, Any] | None,
    capture_mode: str = "auto",
) -> dict[str, Any]:
    should_capture, reason = should_capture_qa(question, answer, mode, capture_mode)
    root = wiki_root(profile)
    if not should_capture:
        return {"saved": False, "reason": reason, "wiki_dir": str(root)}

    questions_dir = wiki_questions_dir(profile)
    syntheses_dir = wiki_syntheses_dir(profile)
    questions_dir.mkdir(parents=True, exist_ok=True)
    syntheses_dir.mkdir(parents=True, exist_ok=True)
    note_name = wiki_note_name(question)
    question_path = questions_dir / note_name
    synthesis_path = syntheses_dir / note_name
    query_context = query_context or {}
    context_path = query_context.get("path")
    now = utc_now()

    question_doc = [
        f"# Question: {question[:120]}",
        "",
        "## Metadata",
        "",
        f"- Profile: `{profile}`",
        f"- Mode: `{mode}`",
        f"- Captured at: `{now}`",
        f"- Capture reason: `{reason}`",
        f"- Query context: `{context_path}`" if context_path else "- Query context: none",
        "",
        "## Question",
        "",
        question.strip(),
        "",
        "## Answer",
        "",
        redact_secrets(answer.strip()),
        "",
        "## Sources",
        "",
        render_sources_for_markdown(sources),
    ]
    question_path.write_text("\n".join(question_doc).rstrip() + "\n", encoding="utf-8")

    synthesis_doc = [
        f"# Synthesis: {question[:120]}",
        "",
        "## Takeaway",
        "",
        redact_secrets(answer.strip()[:6000]),
        "",
        "## Provenance",
        "",
        f"- Derived from question page: [{question_path.name}]({relative_link(synthesis_path, question_path)})",
        f"- Profile: `{profile}`",
        f"- Mode: `{mode}`",
        f"- Captured at: `{now}`",
        "",
        "## Maintenance Notes",
        "",
        "- Treat this as accumulated working knowledge, not a replacement for source retrieval.",
        "- Re-run the question or refresh the folder when source files change materially.",
    ]
    synthesis_path.write_text("\n".join(synthesis_doc).rstrip() + "\n", encoding="utf-8")

    append_wiki_index_item(questions_dir / "index.md", "Saved Questions", "High-value DataMind questions captured from Codex conversations.", question_path)
    append_wiki_index_item(syntheses_dir / "index.md", "Syntheses", "Reusable answer-derived synthesis notes.", synthesis_path)
    append_markdown_log(root / "log.md", f"- {now}: Captured QA to Wiki for profile `{profile}`: [{question_path.stem}]({relative_link(root / 'log.md', question_path)}).")

    return {
        "saved": True,
        "reason": reason,
        "question_path": str(question_path),
        "synthesis_path": str(synthesis_path),
        "wiki_dir": str(root),
    }


def extract_markdown_links(text: str) -> list[str]:
    links = []
    for match in re.finditer(r"\[[^\]]+\]\(([^)]+)\)", text):
        target = match.group(1).strip()
        if target and not re.match(r"^[a-zA-Z][a-zA-Z0-9+.-]*:", target):
            links.append(target.split("#", 1)[0])
    return links


def lint_wiki(args: dict[str, Any]) -> dict[str, Any]:
    profile, _folder = resolve_profile_args(args)
    root = wiki_root(profile)
    issues: list[dict[str, Any]] = []
    required = [
        root / "index.md",
        root / "log.md",
        root / "schema.md",
        root / "user" / "preferences.md",
    ]
    for path in required:
        if not path.is_file():
            issues.append({"severity": "error", "kind": "missing_required_page", "path": str(path)})

    markdown_files = sorted(root.rglob("*.md")) if root.is_dir() else []
    if not markdown_files:
        issues.append({"severity": "error", "kind": "empty_wiki", "path": str(root)})

    for path in markdown_files:
        text = path.read_text(encoding="utf-8", errors="replace")
        if not text.strip():
            issues.append({"severity": "warning", "kind": "empty_page", "path": str(path)})
        if len(text.strip()) < 40:
            issues.append({"severity": "info", "kind": "very_short_page", "path": str(path)})
        for link in extract_markdown_links(text):
            target = (path.parent / link).resolve()
            if not target.exists():
                issues.append({"severity": "warning", "kind": "broken_link", "path": str(path), "target": link})

    source_pages = sorted((root / "sources").glob("*.md")) if (root / "sources").is_dir() else []
    index_text = read_text_file(root / "index.md", 200000)
    for source_page in source_pages:
        if source_page.name not in index_text and relative_link(root / "index.md", source_page) not in index_text:
            issues.append({"severity": "info", "kind": "source_not_linked_from_index", "path": str(source_page)})

    question_pages = [p for p in (root / "questions").glob("*.md") if p.name != "index.md"] if (root / "questions").is_dir() else []
    synthesis_pages = [p for p in (root / "syntheses").glob("*.md") if p.name != "index.md"] if (root / "syntheses").is_dir() else []
    if question_pages and not synthesis_pages:
        issues.append({"severity": "warning", "kind": "questions_without_syntheses", "count": len(question_pages)})

    counts = {
        "errors": sum(1 for issue in issues if issue["severity"] == "error"),
        "warnings": sum(1 for issue in issues if issue["severity"] == "warning"),
        "info": sum(1 for issue in issues if issue["severity"] == "info"),
        "pages": len(markdown_files),
        "sources": len(source_pages),
        "questions": len(question_pages),
        "syntheses": len(synthesis_pages),
    }
    recommendations = []
    if counts["errors"]:
        recommendations.append("Run datamind_use_folder or datamind_ingest_folder to regenerate the profile Wiki.")
    if counts["warnings"]:
        recommendations.append("Review broken links or missing synthesis pages before relying on this Wiki for long-running work.")
    if not recommendations:
        recommendations.append("Wiki looks healthy for the current MVP checks.")

    result = {
        "profile": profile,
        "wiki_dir": str(root),
        "counts": counts,
        "issues": issues[:100],
        "recommendations": recommendations,
    }
    if args.get("write_report", True) is not False:
        report_path = root / "lint.md"
        report_lines = [
            f"# Wiki Lint: {profile}",
            "",
            f"Generated at: `{utc_now()}`",
            "",
            "## Counts",
            "",
            *(f"- {key}: `{value}`" for key, value in counts.items()),
            "",
            "## Issues",
            "",
        ]
        if issues:
            for issue in issues[:100]:
                report_lines.append(f"- `{issue['severity']}` `{issue['kind']}`: {json.dumps(issue, ensure_ascii=False)}")
        else:
            report_lines.append("- No issues found.")
        report_lines.extend(["", "## Recommendations", ""])
        report_lines.extend(f"- {item}" for item in recommendations)
        report_path.parent.mkdir(parents=True, exist_ok=True)
        report_path.write_text("\n".join(report_lines).rstrip() + "\n", encoding="utf-8")
        result["report_path"] = str(report_path)
    return tool_result(result)


def refresh_folder(args: dict[str, Any]) -> dict[str, Any]:
    profile, folder = resolve_profile_args(args, require_folder=True)
    assert folder is not None
    payload = {
        "profile": profile,
        "folder_path": str(folder),
        "rebuild_rag": args.get("rebuild_rag", True),
        "rebuild_graph": args.get("rebuild_graph", True),
        "force_rebuild": args.get("force_rebuild", True),
        "similarity_top_k": args.get("similarity_top_k"),
        "timeout_seconds": args.get("timeout_seconds", DEFAULT_TIMEOUT_SECONDS),
    }
    output = run_task("refresh", payload)
    entry = remember_profile(profile, folder)
    output["registered_profile"] = entry
    manifest_path = MANAGED_PROFILES_DIR / profile / "manifest.json"
    if manifest_path.is_file() and args.get("build_wiki", True) is not False:
        try:
            manifest = json.loads(manifest_path.read_text(encoding="utf-8"))
            output["wiki"] = write_profile_wiki(profile, manifest, entry)
        except Exception as exc:
            output["wiki_error"] = redact_secrets(f"{type(exc).__name__}: {exc}")
    return tool_result(output)


def ingest_folder(args: dict[str, Any]) -> dict[str, Any]:
    profile, source, data_dir, manifest = materialize_folder(args)
    entry = remember_profile(profile, data_dir, source_folder=source, managed=True)
    refresh = args.get("refresh", True) is not False
    changes = manifest.get("changes", {}) if isinstance(manifest, dict) else {}
    indexed_change_count = sum(
        len(changes.get(key, []) or [])
        for key in ("indexed_added", "indexed_changed", "indexed_removed")
    )
    output: dict[str, Any] = {
        "profile": profile,
        "source_folder": str(source),
        "data_dir": str(data_dir),
        "manifest": manifest,
        "registered_profile": entry,
    }
    if args.get("build_wiki", True) is not False:
        output["wiki"] = write_profile_wiki(profile, manifest, entry)
    if refresh:
        requested_force_rebuild = args.get("force_rebuild", True) is True
        effective_force_rebuild = requested_force_rebuild or indexed_change_count > 0
        if indexed_change_count == 0 and not requested_force_rebuild:
            status = run_task(
                "status",
                {
                    "profile": profile,
                    "folder_path": str(data_dir),
                    "timeout_seconds": args.get("timeout_seconds", DEFAULT_TIMEOUT_SECONDS),
                },
            )
            vector_counts = status.get("vector_counts", {}) or {}
            has_rag = any(int(value or 0) > 0 for value in vector_counts.values())
            has_graph = status.get("graph_dir_exists") is True
            needs_rag = args.get("rebuild_rag", True) is not False and not has_rag
            needs_graph = args.get("rebuild_graph", True) is not False and not has_graph
            if not needs_rag and not needs_graph:
                output["refresh"] = {
                    "skipped": True,
                    "reason": "no index-relevant file changes detected",
                    "vector_counts": vector_counts,
                    "graph_dir_exists": has_graph,
                }
                return tool_result(output)

        payload = {
            "profile": profile,
            "folder_path": str(data_dir),
            "rebuild_rag": args.get("rebuild_rag", True),
            "rebuild_graph": args.get("rebuild_graph", True),
            "force_rebuild": effective_force_rebuild,
            "similarity_top_k": args.get("similarity_top_k"),
            "timeout_seconds": args.get("timeout_seconds", DEFAULT_TIMEOUT_SECONDS),
        }
        output["refresh"] = run_task("refresh", payload)
    return tool_result(output)


def consolidate_interactions(args: dict[str, Any]) -> dict[str, Any]:
    session_id = default_session_id(args.get("session_id"))
    events_path = interaction_events_path(session_id)
    if not events_path.is_file():
        raise ToolError(f"No recorded interactions for session_id: {session_id}")
    memory_path = memory_file_path(session_id)
    skill_path = skill_file_path(session_id)
    profile, folder = resolve_profile_args(args) if args.get("profile") or args.get("folder_path") else ("default", None)
    payload = {
        "session_id": session_id,
        "profile": profile,
        "folder_path": str(folder) if folder else None,
        "events_path": str(events_path),
        "memory_path": str(memory_path),
        "skill_path": str(skill_path),
        "max_events": args.get("max_events", 80),
        "timeout_seconds": args.get("timeout_seconds", DEFAULT_TIMEOUT_SECONDS),
    }
    output = run_task("consolidate_memory", payload)
    structured_items = []
    memory_doc = read_text_file(memory_path, max_chars=20000)
    if memory_doc:
        item = add_structured_memory_item(
            content=memory_doc,
            scope="session",
            profile=profile if profile != "default" else None,
            session_id=session_id,
            kind="summary",
            title=f"Consolidated memory: {session_id}",
            source="consolidate_memory",
            source_ref=str(memory_path),
            metadata={"events_path": str(events_path), "max_events": payload["max_events"]},
            confidence=0.75,
        )
        if item:
            structured_items.append({"id": item["id"], "kind": item["kind"], "scope": item["scope"]})
    skill_doc = read_text_file(skill_path, max_chars=20000)
    if skill_doc:
        item = add_structured_memory_item(
            content=skill_doc,
            scope="session",
            profile=profile if profile != "default" else None,
            session_id=session_id,
            kind="skill",
            title=f"Generated skill: {session_id}",
            source="consolidate_memory",
            source_ref=str(skill_path),
            metadata={"events_path": str(events_path), "max_events": payload["max_events"]},
            confidence=0.7,
        )
        if item:
            structured_items.append({"id": item["id"], "kind": item["kind"], "scope": item["scope"]})
    output["structured_memory_items"] = structured_items

    inject_profile = args.get("inject_profile") is True
    if inject_profile:
        if folder is None:
            raise ToolError("inject_profile=true requires a managed profile or folder_path.")
        dynamic_dir = folder / "_datamind_dynamic"
        dynamic_dir.mkdir(parents=True, exist_ok=True)
        injected_memory = dynamic_dir / f"{session_slug(session_id)}-memory.md"
        injected_skill = dynamic_dir / f"{session_slug(session_id)}-skill.md"
        shutil.copy2(memory_path, injected_memory)
        shutil.copy2(skill_path, injected_skill)
        output["injected_files"] = [str(injected_memory), str(injected_skill)]
        if args.get("refresh_profile") is True:
            output["refresh"] = run_task(
                "refresh",
                {
                    "profile": profile,
                    "folder_path": str(folder),
                    "rebuild_rag": True,
                    "rebuild_graph": args.get("rebuild_graph", False),
                    "force_rebuild": True,
                    "timeout_seconds": args.get("timeout_seconds", DEFAULT_TIMEOUT_SECONDS),
                },
            )
    return tool_result(output)


def use_folder(args: dict[str, Any]) -> dict[str, Any]:
    """Simple one-step entrypoint: ingest a folder and refresh indexes."""
    folder = expand_path(str(args.get("folder_path") or ""))
    if folder is None or not folder.is_dir():
        raise ToolError(f"folder_path does not exist or is not a directory: {folder}")
    requested_profile = str(args.get("profile") or "").strip()
    profile = normalize_profile(requested_profile or derive_profile(folder))
    profile_existed = profile in load_profiles().get("profiles", {})
    full_args = {
        "folder_path": str(folder),
        "profile": profile,
        "refresh": True,
        "rebuild_rag": True,
        "rebuild_graph": args.get("build_graph", args.get("rebuild_graph", True)),
        "force_rebuild": args.get("force_rebuild", False),
        "clean": args.get("clean", False),
        "max_table_rows": args.get("max_table_rows", 500),
        "transcribe_media": args.get("transcribe_media", False),
        "build_wiki": args.get("build_wiki", True),
        "timeout_seconds": args.get("timeout_seconds", DEFAULT_TIMEOUT_SECONDS),
    }
    output = ingest_folder(full_args)["structuredContent"]
    manifest = output.get("manifest", {})
    refresh = output.get("refresh", {})
    wiki = output.get("wiki", {})
    auto_record_global(
        (
            f"DataMind profile '{profile}' is bound to source folder '{folder}'. "
            f"Managed data lives in '{output.get('data_dir')}'."
        ),
        profile=profile,
        metadata={
            "kind": "profile_binding",
            "source_folder": str(folder),
            "managed_data_dir": output.get("data_dir"),
            "profile_auto_created": not requested_profile,
        },
    )
    next_steps = [
        f"Ask with datamind_ask using profile={profile}, or omit profile to use the last DataMind profile.",
        f"High-value preferences should be auto-recorded into the global session '{GLOBAL_SESSION_ID}'.",
        "Run datamind_use_folder again after adding or changing files.",
    ]
    if wiki.get("index_path"):
        next_steps.insert(1, f"Review the generated Wiki at {wiki.get('index_path')}.")
    return tool_result(
        {
            "profile": profile,
            "profile_was_auto_created": not requested_profile,
            "profile_previously_existed": profile_existed,
            "message": (
                f"Created profile '{profile}' automatically."
                if not requested_profile and not profile_existed
                else f"Using profile '{profile}'."
            ),
            "source_folder": str(folder),
            "managed_data_dir": output.get("data_dir"),
            "wiki": wiki,
            "default_session": GLOBAL_SESSION_ID,
            "file_counts": manifest.get("counts", {}),
            "file_changes": manifest.get("changes", {}),
            "refresh": {
                "skipped": refresh.get("skipped", False),
                "reason": refresh.get("reason"),
                "force_rebuild": refresh.get("force_rebuild"),
            },
            "vector_counts": refresh.get("vector_counts", {}),
            "graph": refresh.get("graph"),
            "next": next_steps,
        }
    )


def is_graph_question(question: str) -> bool:
    keywords = [
        "关系", "关联", "依赖", "路径", "链路", "多跳", "图谱", "连接",
        "包含", "基于", "属于", "影响", "上下游", "relation", "graph",
        "dependency", "depends", "connect", "linked",
    ]
    lower = question.lower()
    return any(keyword in lower for keyword in keywords)


def ask(args: dict[str, Any]) -> dict[str, Any]:
    """Simple query entrypoint over static knowledge plus optional dynamic memory."""
    question = str(args.get("question") or args.get("query") or "").strip()
    if not question:
        raise ToolError("question is required.")
    mode = str(args.get("mode") or "auto").lower()
    wiki_capture = str(args.get("wiki_capture") or "auto").lower()
    profile = args.get("profile")
    folder_path = args.get("folder_path")
    session_id = default_session_id(args.get("session_id"))
    include_memory = args.get("include_memory", True)
    timeout_seconds = args.get("timeout_seconds", DEFAULT_TIMEOUT_SECONDS)

    base_args = {
        "query": question,
        "profile": profile,
        "folder_path": folder_path,
        "session_id": session_id,
        "include_memory": include_memory,
        "similarity_top_k": args.get("similarity_top_k"),
        "timeout_seconds": timeout_seconds,
    }

    if mode == "auto":
        mode = "graph" if is_graph_question(question) else "rag"

    if mode == "rag":
        output = rag_query(base_args)["structuredContent"]
        output["mode_used"] = "rag"
        output["wiki_capture"] = save_qa_to_wiki(
            profile=output.get("profile") or resolve_profile_args(base_args)[0],
            question=question,
            answer=str(output.get("answer") or ""),
            mode="rag",
            sources=output.get("sources") or [],
            query_context=output.get("query_context"),
            capture_mode=wiki_capture,
        )
        return tool_result(output)
    if mode == "graph":
        output = graph_query(base_args)["structuredContent"]
        output["mode_used"] = "graph"
        output["wiki_capture"] = save_qa_to_wiki(
            profile=output.get("profile") or resolve_profile_args(base_args)[0],
            question=question,
            answer=str(output.get("answer") or ""),
            mode="graph",
            sources=output.get("sources") or [],
            query_context=output.get("query_context"),
            capture_mode=wiki_capture,
        )
        return tool_result(output)
    if mode == "both":
        rag_output = rag_query(base_args)["structuredContent"]
        graph_output = graph_query(base_args)["structuredContent"]
        answer = "\n\n## RAG Answer\n\n{rag}\n\n## Graph Answer\n\n{graph}".format(
            rag=rag_output.get("answer") or "",
            graph=graph_output.get("answer") or "",
        )
        profile_used = rag_output.get("profile") or graph_output.get("profile") or resolve_profile_args(base_args)[0]
        capture = save_qa_to_wiki(
            profile=profile_used,
            question=question,
            answer=answer,
            mode="both",
            sources=(rag_output.get("sources") or []) + (graph_output.get("sources") or []),
            query_context=rag_output.get("query_context") or graph_output.get("query_context"),
            capture_mode=wiki_capture,
        )
        feedback_trace = record_answer_trace(
            profile=profile_used,
            session_id=session_id,
            mode="both",
            query=question,
            answer=answer,
            sources=(rag_output.get("sources") or []) + (graph_output.get("sources") or []),
            query_context=rag_output.get("query_context") or graph_output.get("query_context") or {},
            metadata={"tool": "datamind_ask", "subtraces": [rag_output.get("feedback_trace"), graph_output.get("feedback_trace")]},
        )
        return tool_result(
            {
                "query": question,
                "mode_used": "both",
                "rag_answer": rag_output.get("answer"),
                "graph_answer": graph_output.get("answer"),
                "rag_sources": rag_output.get("sources", []),
                "graph_sources": graph_output.get("sources", []),
                "profile": profile_used,
                "memory_context_used": bool(rag_output.get("memory_context_used") or graph_output.get("memory_context_used")),
                "query_context": rag_output.get("query_context") or graph_output.get("query_context"),
                "wiki_capture": capture,
                "feedback_trace": feedback_trace,
            }
        )
    raise ToolError("mode must be one of: auto, rag, graph, both.")


def remember(args: dict[str, Any]) -> dict[str, Any]:
    """Simple memory entrypoint for one note or a short interaction."""
    content = str(args.get("content") or args.get("note") or "").strip()
    messages = args.get("messages")
    if not content and not messages:
        raise ToolError("content or messages is required.")
    record_args = {
        "session_id": default_session_id(args.get("session_id")),
        "profile": args.get("profile"),
        "source": args.get("source", "codex"),
        "metadata": args.get("metadata") or {"kind": "manual_memory"},
    }
    if messages:
        record_args["messages"] = messages
    else:
        record_args["messages"] = [{"role": args.get("role", "user"), "content": content}]
    output = record_interaction(record_args)["structuredContent"]
    preference_text = content
    if not preference_text and isinstance(messages, list):
        parts = []
        for message in messages:
            if isinstance(message, dict) and str(message.get("content") or "").strip():
                parts.append(str(message.get("content")).strip())
        preference_text = "\n".join(parts)
    structured_item = add_structured_note(
        preference_text,
        profile=args.get("profile"),
        session_id=record_args["session_id"],
        metadata=record_args["metadata"],
        source=record_args["source"],
        source_ref=output.get("events_path", ""),
    )
    output["wiki_preferences"] = append_preference_to_wiki(
        preference_text,
        profile=args.get("profile"),
        metadata=record_args["metadata"],
    )
    output["structured_memory"] = (
        {
            "id": structured_item.get("id"),
            "kind": structured_item.get("kind"),
            "scope": structured_item.get("scope"),
            "path": str(MEMORY_DB_PATH),
        }
        if structured_item
        else {"path": str(MEMORY_DB_PATH), "written": False}
    )
    output["next"] = f"Run datamind_save_memory to summarize the global session '{GLOBAL_SESSION_ID}' into Memory and Skill files."
    return tool_result(output)


def save_memory(args: dict[str, Any]) -> dict[str, Any]:
    """Simple one-step consolidation entrypoint."""
    inject = args.get("inject_profile")
    if inject is None:
        inject = bool(args.get("profile") or args.get("folder_path"))
    full_args = {
        "session_id": default_session_id(args.get("session_id")),
        "profile": args.get("profile"),
        "folder_path": args.get("folder_path"),
        "max_events": args.get("max_events", 80),
        "inject_profile": inject,
        "refresh_profile": args.get("refresh_profile", inject),
        "rebuild_graph": args.get("rebuild_graph", False),
        "timeout_seconds": args.get("timeout_seconds", DEFAULT_TIMEOUT_SECONDS),
    }
    output = consolidate_interactions(full_args)["structuredContent"]
    output["next"] = f"Ask with datamind_ask; no session_id is needed because '{GLOBAL_SESSION_ID}' is the default session."
    return tool_result(output)


def rag_query(args: dict[str, Any]) -> dict[str, Any]:
    profile, folder = resolve_profile_args(args)
    session_id = default_session_id(args.get("session_id")) if args.get("include_memory", True) is not False else None
    context_text, context_path, context_meta = write_query_context(
        profile,
        session_id,
        include_memory=args.get("include_memory", True) is not False,
        query=str(args.get("query") or ""),
    )
    payload = {
        "profile": profile,
        "folder_path": str(folder) if folder else None,
        "query": args.get("query"),
        "similarity_top_k": args.get("similarity_top_k"),
        "memory_context_path": context_path if context_text else None,
        "timeout_seconds": args.get("timeout_seconds", DEFAULT_TIMEOUT_SECONDS),
    }
    output = run_task("rag_query", payload)
    if context_path:
        output["query_context_path"] = context_path
        output["query_context"] = context_meta
    if session_id:
        output["session_id"] = session_id
    if folder is not None:
        output["registered_profile"] = remember_profile(profile, folder)
    output["feedback_trace"] = record_answer_trace(
        profile=output.get("profile") or profile,
        session_id=session_id,
        mode="rag",
        query=str(args.get("query") or ""),
        answer=str(output.get("answer") or ""),
        sources=output.get("sources") or [],
        query_context=context_meta if context_path else {},
        metadata={"tool": "datamind_rag_query"},
    )
    return tool_result(output)


def graph_query(args: dict[str, Any]) -> dict[str, Any]:
    profile, folder = resolve_profile_args(args)
    session_id = default_session_id(args.get("session_id")) if args.get("include_memory", True) is not False else None
    context_text, context_path, context_meta = write_query_context(
        profile,
        session_id,
        include_memory=args.get("include_memory", True) is not False,
        query=str(args.get("query") or ""),
    )
    payload = {
        "profile": profile,
        "folder_path": str(folder) if folder else None,
        "query": args.get("query"),
        "memory_context_path": context_path if context_text else None,
        "timeout_seconds": args.get("timeout_seconds", DEFAULT_TIMEOUT_SECONDS),
    }
    output = run_task("graph_query", payload)
    if context_path:
        output["query_context_path"] = context_path
        output["query_context"] = context_meta
    if session_id:
        output["session_id"] = session_id
    if folder is not None:
        output["registered_profile"] = remember_profile(profile, folder)
    output["feedback_trace"] = record_answer_trace(
        profile=output.get("profile") or profile,
        session_id=session_id,
        mode="graph",
        query=str(args.get("query") or ""),
        answer=str(output.get("answer") or ""),
        sources=output.get("sources") or [],
        query_context=context_meta if context_path else {},
        metadata={"tool": "datamind_graph_query"},
    )
    return tool_result(output)


def list_profiles(_: dict[str, Any]) -> dict[str, Any]:
    payload = load_profiles()
    return tool_result(
        {
            "profiles": sorted(payload.get("profiles", {}).values(), key=lambda item: item.get("profile", "")),
            "last_profile": payload.get("last_profile"),
            "default_session": GLOBAL_SESSION_ID,
            "config_path": str(PROFILES_PATH),
        }
    )


def status(args: dict[str, Any]) -> dict[str, Any]:
    profile, folder = resolve_profile_args(args)
    payload = {
        "profile": profile,
        "folder_path": str(folder) if folder else None,
        "timeout_seconds": args.get("timeout_seconds", 120),
    }
    output = run_task("status", payload, timeout_seconds=int(payload["timeout_seconds"]))
    if folder is not None:
        output["registered_profile"] = remember_profile(profile, folder)
    root = wiki_root(profile)
    output["wiki"] = {
        "wiki_dir": str(root),
        "exists": root.is_dir(),
        "index_path": str(root / "index.md"),
        "index_exists": (root / "index.md").is_file(),
        "questions_dir": str(root / "questions"),
        "syntheses_dir": str(root / "syntheses"),
        "question_count": len([p for p in (root / "questions").glob("*.md") if p.name != "index.md"]) if (root / "questions").is_dir() else 0,
        "synthesis_count": len([p for p in (root / "syntheses").glob("*.md") if p.name != "index.md"]) if (root / "syntheses").is_dir() else 0,
        "global_preferences_path": str(global_preferences_path()),
    }
    conn = None
    try:
        conn = memory_store.connect(MEMORY_DB_PATH)
        output["feedback"] = memory_store.feedback_stats(conn, profile=profile, session_id=GLOBAL_SESSION_ID)
    except Exception as exc:
        output["feedback"] = {"path": str(MEMORY_DB_PATH), "error": str(exc)}
    finally:
        if conn is not None:
            conn.close()
    return tool_result(output)


TOOLS: dict[str, dict[str, Any]] = {
    "datamind_use_folder": {
        "description": "Simple entrypoint: ingest a source folder into a managed profile, refresh DataMind indexes, and generate a lightweight Markdown Wiki. If profile is omitted, DataMind creates one automatically and makes it the last-used profile.",
        "inputSchema": {
            "type": "object",
            "required": ["folder_path"],
            "properties": {
                "folder_path": {"type": "string", "description": "Source folder to use as a DataMind knowledge base."},
                "profile": {"type": "string", "description": "Optional short name to reuse later. If omitted, DataMind creates one from the folder path."},
                "build_graph": {"type": "boolean", "default": True},
                "clean": {"type": "boolean", "default": False, "description": "Clear the managed profile data folder before extracting. The simple entrypoint defaults to incremental change detection."},
                "force_rebuild": {"type": "boolean", "default": False, "description": "Force RAG/GraphRAG rebuild even when no index-relevant file changes are detected."},
                "max_table_rows": {"type": "integer", "minimum": 1, "maximum": 5000, "default": 500},
                "transcribe_media": {"type": "boolean", "default": False},
                "build_wiki": {"type": "boolean", "default": True, "description": "Generate or refresh the lightweight Markdown Wiki layer."},
                "timeout_seconds": {"type": "integer", "minimum": 30, "maximum": 7200, "default": DEFAULT_TIMEOUT_SECONDS},
            },
            "additionalProperties": False,
        },
        "handler": use_folder,
    },
    "datamind_ask": {
        "description": "Simple entrypoint: ask DataMind using the last profile by default, Wiki context plus global session memory by default, automatic RAG/GraphRAG selection, and high-value QA capture into the Wiki.",
        "inputSchema": {
            "type": "object",
            "required": ["question"],
            "properties": {
                "question": {"type": "string"},
                "profile": {"type": "string"},
                "folder_path": {"type": "string"},
                "session_id": {"type": "string", "default": GLOBAL_SESSION_ID},
                "include_memory": {"type": "boolean", "default": True},
                "mode": {"type": "string", "enum": ["auto", "rag", "graph", "both"], "default": "auto"},
                "wiki_capture": {"type": "string", "enum": ["auto", "always", "never"], "default": "auto", "description": "Whether to save the question and answer into the profile Wiki."},
                "similarity_top_k": {"type": "integer", "minimum": 1, "maximum": 20},
                "timeout_seconds": {"type": "integer", "minimum": 30, "maximum": 7200, "default": DEFAULT_TIMEOUT_SECONDS},
            },
            "additionalProperties": False,
        },
        "handler": ask,
    },
    "datamind_remember": {
        "description": "Simple entrypoint: remember a user preference, decision, or high-value interaction in the global session and Wiki preferences by default.",
        "inputSchema": {
            "type": "object",
            "properties": {
                "session_id": {"type": "string", "default": GLOBAL_SESSION_ID},
                "profile": {"type": "string"},
                "content": {"type": "string"},
                "note": {"type": "string"},
                "role": {"type": "string", "default": "user"},
                "messages": {
                    "type": "array",
                    "items": {
                        "type": "object",
                        "properties": {
                            "role": {"type": "string"},
                            "content": {"type": "string"},
                            "timestamp": {"type": "string"},
                            "metadata": {"type": "object"},
                        },
                        "required": ["content"],
                        "additionalProperties": True,
                    },
                },
                "metadata": {"type": "object"},
            },
            "additionalProperties": False,
        },
        "handler": remember,
    },
    "datamind_save_memory": {
        "description": "Simple entrypoint: consolidate recorded interactions into Memory and Skill files, optionally injecting them into a profile.",
        "inputSchema": {
            "type": "object",
            "properties": {
                "session_id": {"type": "string", "default": GLOBAL_SESSION_ID},
                "profile": {"type": "string"},
                "folder_path": {"type": "string"},
                "max_events": {"type": "integer", "minimum": 1, "maximum": 500, "default": 80},
                "inject_profile": {"type": "boolean"},
                "refresh_profile": {"type": "boolean"},
                "rebuild_graph": {"type": "boolean", "default": False},
                "timeout_seconds": {"type": "integer", "minimum": 30, "maximum": 7200, "default": DEFAULT_TIMEOUT_SECONDS},
            },
            "additionalProperties": False,
        },
        "handler": save_memory,
    },
    "datamind_ingest_folder": {
        "description": "Materialize a source folder into a managed DataMind profile, extracting common office files, generating a lightweight Wiki, then refreshing RAG/GraphRAG indexes.",
        "inputSchema": {
            "type": "object",
            "required": ["folder_path"],
            "properties": {
                "folder_path": {"type": "string", "description": "Source folder containing static files."},
                "profile": {"type": "string", "description": "Stable profile name. Defaults to a slug derived from the folder name."},
                "refresh": {"type": "boolean", "default": True},
                "rebuild_rag": {"type": "boolean", "default": True},
                "rebuild_graph": {"type": "boolean", "default": True},
                "force_rebuild": {"type": "boolean", "default": True},
                "clean": {"type": "boolean", "default": True, "description": "Clear the managed profile data folder before extracting."},
                "max_table_rows": {"type": "integer", "minimum": 1, "maximum": 5000, "default": 500},
                "transcribe_media": {"type": "boolean", "default": False, "description": "Best-effort transcription for audio/video files through the configured OpenAI-compatible API."},
                "build_wiki": {"type": "boolean", "default": True, "description": "Generate or refresh the lightweight Markdown Wiki layer."},
                "similarity_top_k": {"type": "integer", "minimum": 1, "maximum": 20},
                "timeout_seconds": {"type": "integer", "minimum": 30, "maximum": 7200, "default": DEFAULT_TIMEOUT_SECONDS},
            },
            "additionalProperties": False,
        },
        "handler": ingest_folder,
    },
    "datamind_refresh_folder": {
        "description": "Register a local folder and rebuild DataMind RAG and/or GraphRAG indexes for it.",
        "inputSchema": {
            "type": "object",
            "required": ["folder_path"],
            "properties": {
                "folder_path": {"type": "string", "description": "Absolute or user-relative folder path to index."},
                "profile": {"type": "string", "description": "Optional stable profile name. Defaults to a slug derived from the folder name."},
                "rebuild_rag": {"type": "boolean", "default": True},
                "rebuild_graph": {"type": "boolean", "default": True},
                "force_rebuild": {"type": "boolean", "default": True},
                "build_wiki": {"type": "boolean", "default": True, "description": "Refresh the lightweight Markdown Wiki layer when a managed manifest exists."},
                "similarity_top_k": {"type": "integer", "minimum": 1, "maximum": 20},
                "timeout_seconds": {"type": "integer", "minimum": 30, "maximum": 7200, "default": DEFAULT_TIMEOUT_SECONDS},
            },
            "additionalProperties": False,
        },
        "handler": refresh_folder,
    },
    "datamind_rag_query": {
        "description": "Ask a question over a DataMind folder/profile using vector RAG. Uses the last profile and Wiki plus global session memory by default.",
        "inputSchema": {
            "type": "object",
            "required": ["query"],
            "properties": {
                "query": {"type": "string"},
                "profile": {"type": "string"},
                "folder_path": {"type": "string", "description": "Optional folder path; remembered as a profile when supplied."},
                "session_id": {"type": "string", "default": GLOBAL_SESSION_ID, "description": "Optional dynamic memory session to include. Defaults to the global session."},
                "include_memory": {"type": "boolean", "default": True},
                "similarity_top_k": {"type": "integer", "minimum": 1, "maximum": 20},
                "timeout_seconds": {"type": "integer", "minimum": 30, "maximum": 7200, "default": DEFAULT_TIMEOUT_SECONDS},
            },
            "additionalProperties": False,
        },
        "handler": rag_query,
    },
    "datamind_graph_query": {
        "description": "Ask a relationship or multi-hop question over a DataMind GraphRAG profile. Uses the last profile and Wiki plus global session memory by default.",
        "inputSchema": {
            "type": "object",
            "required": ["query"],
            "properties": {
                "query": {"type": "string"},
                "profile": {"type": "string"},
                "folder_path": {"type": "string", "description": "Optional folder path; remembered as a profile when supplied."},
                "session_id": {"type": "string", "default": GLOBAL_SESSION_ID, "description": "Optional dynamic memory session to include. Defaults to the global session."},
                "include_memory": {"type": "boolean", "default": True},
                "timeout_seconds": {"type": "integer", "minimum": 30, "maximum": 7200, "default": DEFAULT_TIMEOUT_SECONDS},
            },
            "additionalProperties": False,
        },
        "handler": graph_query,
    },
    "datamind_record_interaction": {
        "description": "Append Codex/user interaction messages to a dynamic session log for later memory and skill consolidation.",
        "inputSchema": {
            "type": "object",
            "properties": {
                "session_id": {"type": "string", "default": GLOBAL_SESSION_ID},
                "profile": {"type": "string"},
                "source": {"type": "string", "default": "codex"},
                "role": {"type": "string", "description": "Single-message role when messages is omitted."},
                "content": {"type": "string", "description": "Single-message content when messages is omitted."},
                "messages": {
                    "type": "array",
                    "items": {
                        "type": "object",
                        "properties": {
                            "role": {"type": "string"},
                            "content": {"type": "string"},
                            "timestamp": {"type": "string"},
                            "metadata": {"type": "object"},
                        },
                        "required": ["content"],
                        "additionalProperties": True,
                    },
                },
                "metadata": {"type": "object"},
            },
            "additionalProperties": False,
        },
        "handler": record_interaction,
    },
    "datamind_consolidate_memory": {
        "description": "Summarize recorded interactions into persistent Memory markdown and a reusable Skill markdown file.",
        "inputSchema": {
            "type": "object",
            "properties": {
                "session_id": {"type": "string", "default": GLOBAL_SESSION_ID},
                "profile": {"type": "string", "description": "Optional profile whose managed data folder should receive the generated memory/skill files."},
                "folder_path": {"type": "string", "description": "Optional managed DataMind data folder."},
                "max_events": {"type": "integer", "minimum": 1, "maximum": 500, "default": 80},
                "inject_profile": {"type": "boolean", "default": False},
                "refresh_profile": {"type": "boolean", "default": False},
                "rebuild_graph": {"type": "boolean", "default": False},
                "timeout_seconds": {"type": "integer", "minimum": 30, "maximum": 7200, "default": DEFAULT_TIMEOUT_SECONDS},
            },
            "additionalProperties": False,
        },
        "handler": consolidate_interactions,
    },
    "datamind_list_memories": {
        "description": "List recorded dynamic interaction sessions, consolidated memories, and generated skill files.",
        "inputSchema": {"type": "object", "properties": {}, "additionalProperties": False},
        "handler": list_memories,
    },
    "datamind_search_memory": {
        "description": "Search the structured DataMind memory store added by the MemOS-inspired MVP.",
        "inputSchema": {
            "type": "object",
            "required": ["query"],
            "properties": {
                "query": {"type": "string"},
                "profile": {"type": "string"},
                "folder_path": {"type": "string"},
                "session_id": {"type": "string", "default": GLOBAL_SESSION_ID},
                "include_session": {"type": "boolean", "default": True},
                "limit": {"type": "integer", "minimum": 1, "maximum": 50, "default": 10},
            },
            "additionalProperties": False,
        },
        "handler": search_structured_memory,
    },
    "datamind_record_feedback": {
        "description": "Record next-state feedback for a recent DataMind answer. Negative feedback is converted into a reusable structured-memory hint.",
        "inputSchema": {
            "type": "object",
            "properties": {
                "trace_id": {"type": "string", "description": "Optional feedback_trace.id returned by datamind_ask, datamind_rag_query, or datamind_graph_query. Defaults to the latest matching trace."},
                "profile": {"type": "string"},
                "folder_path": {"type": "string"},
                "session_id": {"type": "string", "default": GLOBAL_SESSION_ID},
                "include_session": {"type": "boolean", "default": True},
                "feedback": {"type": "string", "description": "The user correction, confirmation, or next-state signal."},
                "next_state": {"type": "string"},
                "content": {"type": "string"},
                "role": {"type": "string", "default": "user"},
                "score": {
                    "oneOf": [{"type": "integer"}, {"type": "string"}],
                    "description": "Optional explicit score: 1 positive, 0 neutral, -1 negative. Strings like positive/negative are also accepted.",
                },
                "label": {"type": "string"},
                "hint": {"type": "string", "description": "Optional explicit reusable hint to save when this feedback should affect future behavior."},
                "source": {"type": "string", "default": "codex"},
            },
            "additionalProperties": False,
        },
        "handler": record_feedback,
    },
    "datamind_feedback_status": {
        "description": "Inspect DataMind feedback traces, next-state feedback events, and recent reusable correction hints.",
        "inputSchema": {
            "type": "object",
            "properties": {
                "profile": {"type": "string"},
                "folder_path": {"type": "string"},
                "session_id": {"type": "string", "default": GLOBAL_SESSION_ID},
                "include_session": {"type": "boolean", "default": True},
                "score": {
                    "oneOf": [{"type": "integer"}, {"type": "string"}],
                    "description": "Optional score filter: 1, 0, -1, or positive/neutral/negative.",
                },
                "limit": {"type": "integer", "minimum": 1, "maximum": 50, "default": 10},
            },
            "additionalProperties": False,
        },
        "handler": feedback_status,
    },
    "datamind_lint_wiki": {
        "description": "Check a profile Wiki for missing required pages, broken links, orphan source pages, and missing syntheses.",
        "inputSchema": {
            "type": "object",
            "properties": {
                "profile": {"type": "string"},
                "folder_path": {"type": "string"},
                "write_report": {"type": "boolean", "default": True},
            },
            "additionalProperties": False,
        },
        "handler": lint_wiki,
    },
    "datamind_list_profiles": {
        "description": "List DataMind folder profiles remembered by this Codex plugin.",
        "inputSchema": {"type": "object", "properties": {}, "additionalProperties": False},
        "handler": list_profiles,
    },
    "datamind_status": {
        "description": "Inspect DataMind profile paths, vector counts, and GraphRAG storage status.",
        "inputSchema": {
            "type": "object",
            "properties": {
                "profile": {"type": "string"},
                "folder_path": {"type": "string"},
                "timeout_seconds": {"type": "integer", "minimum": 30, "maximum": 600, "default": 120},
            },
            "additionalProperties": False,
        },
        "handler": status,
    },
}


def handle_request(message: dict[str, Any]) -> dict[str, Any] | None:
    method = message.get("method")
    request_id = message.get("id")
    params = message.get("params") or {}

    if request_id is None:
        return None

    try:
        if method == "initialize":
            protocol_version = params.get("protocolVersion") or DEFAULT_PROTOCOL_VERSION
            return response(
                request_id,
                {
                    "protocolVersion": protocol_version,
                    "capabilities": {"tools": {}},
                    "serverInfo": {"name": SERVER_NAME, "version": SERVER_VERSION},
                },
            )
        if method == "ping":
            return response(request_id, {})
        if method == "tools/list":
            tools = [
                {
                    "name": name,
                    "description": spec["description"],
                    "inputSchema": spec["inputSchema"],
                }
                for name, spec in TOOLS.items()
            ]
            return response(request_id, {"tools": tools})
        if method == "tools/call":
            tool_name = params.get("name")
            arguments = params.get("arguments") or {}
            if tool_name not in TOOLS:
                raise ToolError(f"Unknown tool: {tool_name}")
            handler: Callable[[dict[str, Any]], dict[str, Any]] = TOOLS[tool_name]["handler"]
            return response(request_id, handler(arguments))
        return error_response(request_id, -32601, f"Method not found: {method}")
    except ToolError as exc:
        return response(request_id, error_result(str(exc)))
    except subprocess.TimeoutExpired:
        return response(request_id, error_result("DataMind task timed out. Try a larger timeout_seconds value."))
    except Exception as exc:
        eprint(traceback.format_exc())
        return response(request_id, error_result(f"Unexpected {type(exc).__name__}: {exc}"))


def response(request_id: Any, result: Any) -> dict[str, Any]:
    return {"jsonrpc": "2.0", "id": request_id, "result": result}


def error_response(request_id: Any, code: int, message: str) -> dict[str, Any]:
    return {"jsonrpc": "2.0", "id": request_id, "error": {"code": code, "message": message}}


def send_message(message: dict[str, Any]) -> None:
    sys.stdout.write(json.dumps(message, ensure_ascii=False, separators=(",", ":")) + "\n")
    sys.stdout.flush()


def run_stdio_server() -> None:
    eprint("MCP stdio server started")
    for line in sys.stdin:
        line = line.strip()
        if not line:
            continue
        try:
            message = json.loads(line)
        except json.JSONDecodeError as exc:
            send_message(error_response(None, -32700, f"Parse error: {exc}"))
            continue
        result = handle_request(message)
        if result is not None:
            send_message(result)


if __name__ == "__main__":
    run_stdio_server()
